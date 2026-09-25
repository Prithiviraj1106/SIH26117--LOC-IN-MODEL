"""
OMEGA AI Workbench - Document Loader
Extracts text from PDF, DOCX, and TXT files completely locally.
- PyMuPDF (fitz) for PDF
- python-docx for DOCX
- Built-in file handling for TXT
"""
import io
from typing import Tuple, Optional


def extract_text_from_pdf(file_bytes: bytes) -> Tuple[str, int]:
    """
    Extracts text from PDF bytes using PyMuPDF (fitz).
    Returns (extracted_text, page_count).
    """
    try:
        import fitz  # PyMuPDF
        
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        page_count = len(doc)
        text_parts = []
        
        for page_num in range(page_count):
            page = doc[page_num]
            page_text = page.get_text("text")
            if page_text and page_text.strip():
                text_parts.append(f"--- Page {page_num + 1} ---\n{page_text.strip()}")
                
        doc.close()
        full_text = "\n\n".join(text_parts)
        return full_text, page_count
        
    except ImportError:
        raise ImportError("PyMuPDF is required for PDF parsing. Please install with `pip install PyMuPDF`.")
    except Exception as e:
        raise RuntimeError(f"Failed to extract text from PDF: {str(e)}")


def extract_text_from_docx(file_bytes: bytes) -> Tuple[str, int]:
    """
    Extracts text from DOCX bytes using python-docx.
    Returns (extracted_text, paragraph_count).
    """
    try:
        import docx
        
        doc_stream = io.BytesIO(file_bytes)
        doc = docx.Document(doc_stream)
        
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        
        # Also extract table contents if any
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                row_str = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_str:
                    table_texts.append(row_str)
                    
        all_parts = paragraphs + table_texts
        full_text = "\n\n".join(all_parts)
        return full_text, len(paragraphs)
        
    except ImportError:
        raise ImportError("python-docx is required for DOCX parsing. Please install with `pip install python-docx`.")
    except Exception as e:
        raise RuntimeError(f"Failed to extract text from DOCX: {str(e)}")


def extract_text_from_txt(file_bytes: bytes) -> Tuple[str, int]:
    """
    Extracts text from raw TXT bytes using standard Python decoding.
    Returns (extracted_text, line_count).
    """
    encodings = ["utf-8", "utf-16", "latin-1", "cp1252"]
    for enc in encodings:
        try:
            text = file_bytes.decode(enc)
            line_count = len(text.splitlines())
            return text, line_count
        except (UnicodeDecodeError, LookupError):
            continue
            
    # Fallback with replacement
    text = file_bytes.decode("utf-8", errors="replace")
    line_count = len(text.splitlines())
    return text, line_count


def load_document(uploaded_file) -> Tuple[str, int, str]:
    """
    Dispatcher to load any supported document from Streamlit's UploadedFile or raw bytes.
    Returns: (extracted_text, page_or_unit_count, file_type)
    """
    if uploaded_file is None:
        return "", 0, "unknown"
        
    filename = getattr(uploaded_file, "name", "document.txt").lower()
    
    # Read bytes
    if hasattr(uploaded_file, "getvalue"):
        file_bytes = uploaded_file.getvalue()
    elif hasattr(uploaded_file, "read"):
        file_bytes = uploaded_file.read()
    else:
        file_bytes = bytes(uploaded_file)
        
    if filename.endswith(".pdf"):
        text, count = extract_text_from_pdf(file_bytes)
        return text, count, "PDF"
    elif filename.endswith(".docx") or filename.endswith(".doc"):
        text, count = extract_text_from_docx(file_bytes)
        return text, count, "DOCX"
    elif filename.endswith(".txt") or filename.endswith(".md"):
        text, count = extract_text_from_txt(file_bytes)
        return text, count, "TXT"
    else:
        # Generic text attempt
        text, count = extract_text_from_txt(file_bytes)
        return text, count, "UNKNOWN"
