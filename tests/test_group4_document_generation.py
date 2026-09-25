import asyncio
from pathlib import Path

import pytest
from fastapi import HTTPException
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

import orchestrator
from api_server import ApprovalRequest, approve_action
from docgen_mcp_server import generate_docx, read_docx
from orchestrator import flow_pptx, flow_word_gen, flow_xlsx, structure_document_text
from policy import check_tool_policy
from pptx_mcp_server import generate_pptx, read_pptx
from xlsx_mcp_server import generate_xlsx, read_xlsx


class FakeManager:
    def __init__(self, results=None):
        self.results = results or {}
        self.calls = []

    async def call_tool(self, name, args):
        self.calls.append((name, args))
        result = self.results.get(name, "")
        return result(args) if callable(result) else result


def test_real_word_generation_and_readback():
    path = Path(generate_docx("Group 4 Word", "Supplied content"))
    assert path.is_file()
    assert "Supplied content" in read_docx(str(path))


def test_extracted_text_is_structured_into_document_blocks():
    structured = structure_document_text(
        "ABOUT ME\n\nThis is a paragraph.\n\nEDUCATION\n\n"
        "Borcelle University | 2026-2030\n\nSKILLS\n\n"
        "- Auditing\n- Financial Accounting\n- Financial Reporting"
    )

    assert structured["blocks"] == [
        {"type": "heading", "text": "ABOUT ME"},
        {"type": "paragraph", "text": "This is a paragraph."},
        {"type": "heading", "text": "EDUCATION"},
        {"type": "paragraph", "text": "Borcelle University | 2026-2030"},
        {"type": "heading", "text": "SKILLS"},
        {"type": "bullet_list", "items": [
            "Auditing", "Financial Accounting", "Financial Reporting",
        ]},
    ]


def test_structured_docx_renders_headings_bullets_and_tables(tmp_path):
    path = Path(generate_docx("Structured", {
        "title": "",
        "blocks": [
            {"type": "heading", "text": "ABOUT ME"},
            {"type": "paragraph", "text": "A paragraph."},
            {"type": "bullet_list", "items": ["Auditing", "Reporting"]},
            {"type": "table", "rows": [["Year", "Role"], ["2026", "Accountant"]]},
        ],
    }))
    document = Document(path)

    assert [(paragraph.text, paragraph.style.name) for paragraph in document.paragraphs] == [
        ("Structured", "Title"),
        ("ABOUT ME", "Heading 1"),
        ("A paragraph.", "Normal"),
        ("Auditing", "List Bullet"),
        ("Reporting", "List Bullet"),
    ]
    assert [[cell.text for cell in row.cells] for row in document.tables[0].rows] == [
        ["Year", "Role"], ["2026", "Accountant"],
    ]


def test_word_generation_accepts_explicit_inline_content(tmp_path):
    artifact = tmp_path / "inline.docx"
    manager = FakeManager({"generate_docx": str(artifact)})

    def write_artifact(args):
        generated = Path(generate_docx(args["title"], args["content"]))
        artifact.write_bytes(generated.read_bytes())
        return str(artifact)

    manager.results["generate_docx"] = write_artifact
    result = asyncio.run(flow_word_gen("Create a Word document: Supplied inline content", None, manager, []))

    assert result == f"Word document generated and saved: {artifact}"
    assert manager.calls[0][1]["content"] == "Supplied inline content"


def test_summary_pdf_word_flow_summarizes_before_docx(monkeypatch, tmp_path):
    artifact = tmp_path / "summary.docx"
    manager = FakeManager({
        "extract_from_document": "PDF source text",
        "generate_docx": str(artifact),
    })

    async def fake_chat(*_args, **_kwargs):
        return "PDF summary text"

    monkeypatch.setattr(orchestrator, "chat_llm", fake_chat)
    monkeypatch.setattr(orchestrator, "make_llm", lambda *_args: object())

    result = asyncio.run(flow_word_gen(
        "summarize the PDF file and then make it as a document",
        str(tmp_path / "source.pdf"),
        manager,
    ))

    assert result.endswith(str(artifact))
    assert [call[0] for call in manager.calls] == ["extract_from_document", "generate_docx"]
    assert manager.calls[1][1]["title"] == "PDF Summary"
    assert manager.calls[1][1]["content"]["blocks"] == [
        {"type": "paragraph", "text": "PDF summary text"},
    ]


def test_prime_numbers_are_generated_deterministically_and_sent_to_docx(tmp_path):
    artifact = tmp_path / "prime-numbers.docx"
    artifact.write_bytes(Path(generate_docx("Temporary", "content")).read_bytes())
    manager = FakeManager({"generate_docx": str(artifact)})

    result = asyncio.run(flow_word_gen(
        "can you give all prime numbers from 1 to 200 in doc",
        None,
        manager,
        [],
    ))

    expected = [
        "2", "3", "5", "7", "11", "13", "17", "19", "23", "29", "31", "37",
        "41", "43", "47", "53", "59", "61", "67", "71", "73", "79", "83", "89",
        "97", "101", "103", "107", "109", "113", "127", "131", "137", "139", "149",
        "151", "157", "163", "167", "173", "179", "181", "191", "193", "197", "199",
    ]
    content = manager.calls[0][1]["content"]
    assert result == f"Word document generated and saved: {artifact}"
    assert content.splitlines() == expected
    assert all(value in content.splitlines() for value in ("2", "3", "197", "199"))
    assert all(value not in content.splitlines() for value in ("1", "4", "200"))


def test_real_excel_generation_and_readback():
    path = Path(generate_xlsx({
        "title": "Group 4 Excel",
        "headers": ["Item", "Total"],
        "rows": [["Pump", 42]],
    }))
    assert path.is_file()
    workbook = load_workbook(path, read_only=True, data_only=True)
    assert list(workbook.active.values)[1] == ("Pump", 42)
    assert "Pump" in read_xlsx(str(path))


def test_real_powerpoint_generation_and_readback():
    path = Path(generate_pptx({
        "title": "Group 4 Presentation",
        "slides_content": "Summary\nTotals are 42",
    }))
    assert path.is_file()
    presentation = Presentation(path)
    assert len(presentation.slides) >= 2
    assert "Totals are 42" in read_pptx(str(path))


def test_incomplete_excel_and_powerpoint_inputs_fail_closed():
    assert generate_xlsx({"title": "Missing rows"}).startswith("Error:")
    assert generate_pptx({"title": "Missing slides"}).startswith("Error:")


def test_scanned_pdf_word_flow_extracts_then_generates(tmp_path):
    artifact = tmp_path / "scan.docx"
    manager = FakeManager({"extract_from_document": "OCR extracted content", "generate_docx": str(artifact)})
    artifact.write_bytes(Path(generate_docx("Temporary", "content")).read_bytes())

    result = asyncio.run(flow_word_gen(
        "Convert this scanned PDF into a Word document",
        str(tmp_path / "scan.pdf"),
        manager,
    ))

    assert result.endswith(str(artifact))
    assert [call[0] for call in manager.calls] == ["extract_from_document", "generate_docx"]
    assert manager.calls[1][1]["content"]["blocks"] == [
        {"type": "paragraph", "text": "OCR extracted content"},
    ]


def test_scanned_report_approval_note_retrieves_authorized_context(tmp_path, monkeypatch):
    artifact = tmp_path / "approval-note.docx"
    artifact.write_bytes(Path(generate_docx("Temporary", "content")).read_bytes())
    manager = FakeManager({
        "extract_from_document": "Finding: pump vibration exceeds the limit.",
        "ingest_document": '{"success": true, "file_hash": "report-hash", "chunks": 1}',
        "search_sops": "[source=pump_maintenance.txt chunk=1]\nInspect vibration weekly.",
        "generate_docx": str(artifact),
    })

    async def fake_chat(*_args, **_kwargs):
        return "Findings\nPump vibration exceeds the limit.\n\nSOP Reference\nWeekly inspection."

    monkeypatch.setattr(orchestrator, "chat_llm", fake_chat)
    monkeypatch.setattr(orchestrator, "make_llm", lambda *_args: object())
    result = asyncio.run(orchestrator.flow_approval_note(
        "Read this scanned inspection report, identify the key findings, and prepare a Word approval note.",
        str(tmp_path / "inspection.pdf"),
        manager,
        owner_id="demo-owner",
    ))

    assert result == f"Approval note drafted and saved: {artifact}"
    assert manager.calls[2] == (
        "search_sops",
        {"query": "Finding: pump vibration exceeds the limit.", "top_k": 5, "owner_id": "demo-owner"},
    )
    assert manager.calls[-1][0] == "generate_docx"


def test_document_flows_propagate_generation_failure_and_approval(monkeypatch):
    async def fake_chat(*_args, **_kwargs):
        return '{"title":"Report","headers":["A"],"rows":"[[1]]"}'

    monkeypatch.setattr(orchestrator, "chat_llm", fake_chat)
    monkeypatch.setattr(orchestrator, "make_llm", lambda *_args: object())
    xlsx_approval = FakeManager({"generate_xlsx": "Approval required: Tool 'generate_xlsx' requires user approval."})
    assert asyncio.run(flow_xlsx("Create an Excel report", xlsx_approval)).startswith("Approval required:")

    async def ppt_chat(*_args, **_kwargs):
        return '{"title":"Report","slides_content":"Summary\\nContent"}'

    monkeypatch.setattr(orchestrator, "chat_llm", ppt_chat)
    ppt_failure = FakeManager({"generate_pptx": "Error: generator failed"})
    assert asyncio.run(flow_pptx("Create a PowerPoint presentation", ppt_failure)).startswith("Presentation validation failed:")


def test_all_document_generators_require_approval_by_default(monkeypatch):
    monkeypatch.delenv("MRPL_LOCAL_DEMO_MODE", raising=False)
    for tool in ("generate_docx", "generate_xlsx", "generate_pptx"):
        assert check_tool_policy(tool).action == "REQUIRE_APPROVAL"
        assert check_tool_policy(tool, {tool}).action == "ALLOW"


def test_api_rejects_unknown_approval_targets():
    with pytest.raises(HTTPException) as error:
        approve_action(ApprovalRequest(session_id="group4", tool="read_xlsx"))
    assert error.value.status_code == 400